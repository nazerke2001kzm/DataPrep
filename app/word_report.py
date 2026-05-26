"""Сборка отчёта Microsoft Word: текст плана, лог, графики, WoE/IV/KS/Gini."""

from __future__ import annotations

import io
import math
from typing import Any

import matplotlib

matplotlib.use("Agg")
import matplotlib.pyplot as plt
import numpy as np
import pandas as pd
from docx import Document
from docx.shared import Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH

from app.binning_metrics import binning_analysis_block, guess_target_column
from app.schemas import AnalystReport


def _fig_to_bytes(fig: plt.Figure, dpi: int = 120) -> bytes:
    buf = io.BytesIO()
    fig.savefig(buf, format="png", dpi=dpi, bbox_inches="tight", facecolor="white")
    plt.close(fig)
    buf.seek(0)
    return buf.read()


def _chart_null_rates(null_rates: dict[str, float], top: int = 15) -> bytes | None:
    items = sorted(null_rates.items(), key=lambda x: -x[1])[:top]
    if not items:
        return None
    names = [str(k)[:40] for k, _ in items]
    vals = [v * 100 for _, v in items]
    fig, ax = plt.subplots(figsize=(8, 4.2))
    ax.barh(names[::-1], vals[::-1], color="#0ea5e9")
    ax.set_xlabel("Доля пропусков, %")
    ax.set_title("Топ колонок по доле NA (исходный файл)")
    ax.set_xlim(0, max(vals) * 1.08 + 1e-6)
    fig.tight_layout()
    return _fig_to_bytes(fig)


def _chart_woe(woe_df: pd.DataFrame, title: str) -> bytes | None:
    if woe_df.empty or "bin" not in woe_df.columns:
        return None
    d = woe_df.sort_values("woe", ascending=True)
    labels = [str(x)[:35] for x in d["bin"]]
    woe = d["woe"].astype(float).values
    fig, ax = plt.subplots(figsize=(8, max(3.0, 0.35 * len(labels))))
    colors = np.where(woe >= 0, "#dc2626", "#16a34a")
    ax.barh(labels, woe, color=colors)
    ax.axvline(0, color="#64748b", linewidth=0.8)
    ax.set_xlabel("WoE")
    ax.set_title(title)
    fig.tight_layout()
    return _fig_to_bytes(fig)


def _chart_histogram(series: pd.Series, title: str) -> bytes | None:
    s = pd.to_numeric(series, errors="coerce").dropna()
    if len(s) < 5:
        return None
    fig, ax = plt.subplots(figsize=(7, 3.8))
    ax.hist(s, bins=min(40, max(10, int(np.sqrt(len(s))))), color="#6366f1", edgecolor="white", alpha=0.9)
    ax.set_title(title)
    ax.set_ylabel("Частота")
    fig.tight_layout()
    return _fig_to_bytes(fig)


def _chart_bin_counts(binned: pd.Series, title: str) -> bytes | None:
    vc = binned.astype("string").value_counts(dropna=True).head(25)
    if vc.empty:
        return None
    fig, ax = plt.subplots(figsize=(8, max(3.2, 0.32 * len(vc))))
    ax.barh(vc.index.astype(str).str.slice(0, 35)[::-1], vc.values[::-1], color="#14b8a6")
    ax.set_xlabel("Количество")
    ax.set_title(title)
    fig.tight_layout()
    return _fig_to_bytes(fig)


def _add_table_from_dataframe(doc: Document, df: pd.DataFrame, max_rows: int = 30) -> None:
    if df.empty:
        return
    show = df.head(max_rows)
    table = doc.add_table(rows=1, cols=len(show.columns))
    table.style = "Table Grid"
    hdr = table.rows[0].cells
    for i, col in enumerate(show.columns):
        hdr[i].text = str(col)
    for _, row in show.iterrows():
        cells = table.add_row().cells
        for i, v in enumerate(row):
            cells[i].text = "" if pd.isna(v) else f"{v:.6g}" if isinstance(v, float) else str(v)


def build_word_report_bytes(
    *,
    filename: str,
    preset_label: str,
    rows_before: int,
    rows_after: int,
    df_original: pd.DataFrame,
    df_prepared: pd.DataFrame,
    report: dict[str, Any],
    plan: AnalystReport,
) -> bytes:
    plt.rcParams["font.family"] = "DejaVu Sans"
    plt.rcParams["axes.unicode_minus"] = False

    doc = Document()
    title = doc.add_heading("Отчёт о подготовке данных", level=0)
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER

    p = doc.add_paragraph()
    p.add_run("Файл: ").bold = True
    p.add_run(filename)
    p = doc.add_paragraph()
    p.add_run("Пресет: ").bold = True
    p.add_run(preset_label)
    p = doc.add_paragraph()
    p.add_run("Строк: ").bold = True
    p.add_run(f"{rows_before} → {rows_after} (после обработки)")

    doc.add_heading("Краткое резюме (план модели)", level=1)
    doc.add_paragraph((report.get("summary_ru") or "—").strip() or "—")

    for label, key in (
        ("Заметки по очистке", "cleaning_notes"),
        ("Заметки по валидации", "validation_notes"),
        ("Заметки по биннингу", "binning_notes"),
        ("Эконометрика (матрица переходов)", "econometric_notes"),
    ):
        notes = report.get(key) or []
        if isinstance(notes, list) and notes:
            doc.add_heading(label, level=2)
            for n in notes:
                doc.add_paragraph(str(n), style="List Bullet")

    doc.add_heading("Журнал выполнения", level=1)
    log = report.get("execution_log_ru") or []
    if isinstance(log, list) and log:
        for line in log:
            doc.add_paragraph(str(line), style="List Bullet")
    else:
        doc.add_paragraph("Шаги не зафиксированы.")

    doc.add_heading("Обзор исходных данных", level=1)
    null_rates = {str(c): float(df_original[c].isna().mean()) for c in df_original.columns}
    doc.add_paragraph(
        f"Размерность: {df_original.shape[0]} строк × {df_original.shape[1]} колонок. "
        f"Ниже — наглядные диаграммы по пропускам и распределениям."
    )
    img = _chart_null_rates(null_rates)
    if img:
        doc.add_picture(io.BytesIO(img), width=Inches(6.2))

    target_col = guess_target_column(df_prepared)
    if target_col:
        doc.add_paragraph(
            f"Обнаружена колонка таргета (эвристика): «{target_col}». "
            "Метрики IV / KS / Gini считаются в предположении, что 1 (или «плохой» класс) — событие риска."
        )
    else:
        doc.add_paragraph(
            "Колонка таргета не распознана по имени (target, default_flag, is_fraud и т.п.). "
            "Таблицы WoE и IV в отчёте не строятся; остаются распределения и счётчики по бинам."
        )

    eco = report.get("transition_econometrics")
    if isinstance(eco, dict) and eco:
        doc.add_heading("Матрица переходов: макрофакторы и эконометрика", level=1)
        tm = report.get("transition_matrix") or {}
        if isinstance(tm, dict):
            doc.add_paragraph(
                f"Корзина: {tm.get('bucket_column') or '—'} | "
                f"Период: {tm.get('period_column') or '—'} | "
                f"Зависимая (adj. R²): {tm.get('dependent_column') or eco.get('dependent_column') or '—'}"
            )
            macros = tm.get("macro_columns") or eco.get("macro_columns") or []
            if macros:
                doc.add_paragraph("Макрофакторы: " + ", ".join(str(m) for m in macros[:15]))

        for line in eco.get("recommendations_ru") or []:
            doc.add_paragraph(str(line), style="List Bullet")

        vif = eco.get("vif") or {}
        if vif.get("ok") and vif.get("rows"):
            doc.add_heading("VIF (мультиколлинеарность)", level=2)
            doc.add_paragraph(vif.get("summary_ru") or "")
            vdf = pd.DataFrame(vif["rows"])
            _add_table_from_dataframe(doc, vdf)

        adj = eco.get("adjusted_r2_ranking") or {}
        if adj.get("ok") and adj.get("ranking"):
            doc.add_heading("Отбор переменных (adjusted R²)", level=2)
            doc.add_paragraph(adj.get("summary_ru") or "")
            rdf = pd.DataFrame(adj["ranking"][:15])
            _add_table_from_dataframe(doc, rdf)

        lags = eco.get("recommended_lags_global") or []
        if lags:
            doc.add_paragraph(f"Рекомендуемые лаги для модели: {lags}")

        stationarity = eco.get("stationarity") or {}
        if stationarity:
            doc.add_heading("Стационарность (ADF / KPSS / PP)", level=2)
            rows = []
            for col, st in list(stationarity.items())[:12]:
                adf = st.get("adf") or {}
                kpss = st.get("kpss") or {}
                pp = st.get("pp") or {}
                rows.append(
                    {
                        "column": col,
                        "ADF_p": adf.get("pvalue"),
                        "KPSS_p": kpss.get("pvalue"),
                        "PP_p": pp.get("pvalue"),
                        "consensus": st.get("consensus_stationary"),
                    }
                )
            _add_table_from_dataframe(doc, pd.DataFrame(rows))

        normality = eco.get("normality") or {}
        if normality:
            doc.add_heading("Нормальность (Jarque–Bera / Shapiro)", level=2)
            nrow = []
            for col, nm in list(normality.items())[:12]:
                jb = (nm.get("jarque_bera") or {}) if isinstance(nm, dict) else {}
                nrow.append(
                    {
                        "column": col,
                        "JB_p": jb.get("pvalue"),
                        "summary": nm.get("summary_ru") if isinstance(nm, dict) else "",
                    }
                )
            _add_table_from_dataframe(doc, pd.DataFrame(nrow))

    doc.add_heading("Анализ биннинга и метрики", level=1)

    if not plan.binning:
        doc.add_paragraph("В плане не указан биннинг признаков.")
    else:
        for spec in plan.binning:
            doc.add_heading(f"Признак: {spec.column} → {spec.output_column}", level=2)
            doc.add_paragraph(
                f"Метод: {spec.method}, число интервалов: {spec.n_bins}. "
                f"Категория в подготовленных данных: «{spec.output_column}»."
            )

            block = binning_analysis_block(
                df_prepared,
                source_col=spec.column,
                binned_col=spec.output_column,
                target_col=target_col,
            )

            wdf = block.get("woe_table")
            if isinstance(wdf, pd.DataFrame) and not wdf.empty:
                doc.add_paragraph(
                    f"IV (Information Value) = {block['iv']:.4f}  |  "
                    f"KS = {block['ks'] * 100:.2f}% (в долях накопления bad/good по бинам, сортировка по WoE)  |  "
                    f"Gini (по непрерывному «{spec.column}») = {block['gini']:.4f}  |  "
                    f"AUC = {block['auc']:.4f}"
                )
                doc.add_paragraph("Таблица WoE по бинам:")
                disp = wdf[["bin", "n", "bad", "good", "bad_rate", "woe", "iv_piece"]].copy()
                _add_table_from_dataframe(doc, disp)
                wimg = _chart_woe(wdf, f"WoE по бинам: {spec.output_column}")
                if wimg:
                    doc.add_picture(io.BytesIO(wimg), width=Inches(6.2))
            else:
                doc.add_paragraph(
                    "WoE / IV / KS не посчитаны (нет подходящего таргета или пустые пересечения после биннинга)."
                )
                if spec.column in df_prepared.columns:
                    auc, gini = block.get("auc", float("nan")), block.get("gini", float("nan"))
                    if math.isfinite(auc):
                        doc.add_paragraph(f"По сырому числовому признаку: AUC = {auc:.4f}, Gini = {gini:.4f}.")

            if spec.column in df_prepared.columns:
                himg = _chart_histogram(
                    df_prepared[spec.column],
                    f"Распределение исходного числового признака «{spec.column}» (после очистки)",
                )
                if himg:
                    doc.add_picture(io.BytesIO(himg), width=Inches(6.0))

            if spec.output_column in df_prepared.columns:
                cimg = _chart_bin_counts(
                    df_prepared[spec.output_column],
                    f"Число наблюдений по бину: {spec.output_column}",
                )
                if cimg:
                    doc.add_picture(io.BytesIO(cimg), width=Inches(6.0))

    doc.add_heading("Приложение: структура плана (JSON)", level=1)
    doc.add_paragraph(
        "Полный машиночитаемый план и параметры очистки/валидации сохранены в файле report.json внутри ZIP-архива."
    )

    out = io.BytesIO()
    doc.save(out)
    return out.getvalue()
